from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class ATSJob:
    job_id: str
    job_name: str
    job_url: str
    best_resume: str
    best_score: float
    recommendation: str
    matched_required_skills: str
    missing_required_skills: str
    evidence_summary: str


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_jobs(
    workbook_path: Path,
    score_threshold: float,
    minimum_score: float = 0.0,
) -> list[ATSJob]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"ATS workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Summary" not in workbook.sheetnames:
        raise ValueError("ATS workbook does not contain a 'Summary' sheet.")

    sheet = workbook["Summary"]
    rows = sheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows)]
    index = {name: position for position, name in enumerate(headers)}

    required = [
        "Job ID", "Job Name", "Job URL", "Best Resume", "Best ATS Score",
        "Recommendation", "Matched Required Skills", "Missing Required Skills",
        "Evidence Summary",
    ]
    missing_columns = [name for name in required if name not in index]
    if missing_columns:
        raise ValueError(f"Missing Summary columns: {', '.join(missing_columns)}")

    selected: list[ATSJob] = []

    for row in rows:
        job_id = _text(row[index["Job ID"]])
        if not job_id:
            continue

        try:
            score = float(row[index["Best ATS Score"]])
        except (TypeError, ValueError):
            continue

        if not (minimum_score <= score < score_threshold):
            continue

        selected.append(
            ATSJob(
                job_id=job_id,
                job_name=_text(row[index["Job Name"]]),
                job_url=_text(row[index["Job URL"]]),
                best_resume=_text(row[index["Best Resume"]]),
                best_score=score,
                recommendation=_text(row[index["Recommendation"]]),
                matched_required_skills=_text(row[index["Matched Required Skills"]]),
                missing_required_skills=_text(row[index["Missing Required Skills"]]),
                evidence_summary=_text(row[index["Evidence Summary"]]),
            )
        )

    return selected
