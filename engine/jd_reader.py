# ==========================================================
# engine/jd_reader.py
# ==========================================================

from __future__ import annotations

import re
from pathlib import Path

from docx import Document


# ==========================================================
# READ DOCUMENT
# ==========================================================

def read_document_text(path: Path) -> str:

    path = Path(path)

    if not path.exists():

        raise FileNotFoundError(
            f"Filtered JD file not found: {path}"
        )

    suffix = path.suffix.lower()

    # ------------------------------------------------------
    # TXT
    # ------------------------------------------------------

    if suffix == ".txt":

        return path.read_text(
            encoding="utf-8",
            errors="replace"
        )

    # ------------------------------------------------------
    # DOCX
    # ------------------------------------------------------

    if suffix == ".docx":

        document = Document(path)

        parts: list[str] = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                parts.append(text)

        # Also support JDs stored inside tables.

        for table in document.tables:

            for row in table.rows:

                values = [

                    cell.text.strip()

                    for cell in row.cells

                    if cell.text.strip()
                ]

                if values:

                    parts.append(
                        " | ".join(values)
                    )

        return "\n".join(parts)

    raise ValueError(
        "Filtered JD input must be .docx or .txt"
    )


# ==========================================================
# NORMALIZE URL
# ==========================================================

def normalize_url(url: str) -> str:

    if not url:

        return ""

    url = str(url).strip()

    # Treat Dice URLs with/without www as identical.

    url = url.replace(
        "https://www.dice.com/",
        "https://dice.com/"
    )

    url = url.replace(
        "http://www.dice.com/",
        "http://dice.com/"
    )

    # Remove query string / fragment.

    url = url.split("?")[0]

    url = url.split("#")[0]

    return url.rstrip("/").lower()


# ==========================================================
# SPLIT JD DOCUMENT
# ==========================================================

def split_job_blocks(text: str) -> list[str]:
    """
    Expected exporter format:

    Job 1: Data Governance Analyst
    Company: ...
    Location: ...
    Job URL: ...
    ...

    Job 2: Data Engineer
    Company: ...
    ...

    Split only on the TOP LEVEL:
        Job <number>: <title>

    This prevents internal lines such as:
        Job 6
        URL:
    from becoming separate JD blocks.
    """

    pattern = r"(?=^Job\s+\d+\s*:\s*.+?$)"

    blocks = re.split(
        pattern,
        text,
        flags=re.MULTILINE
    )

    cleaned_blocks: list[str] = []

    for block in blocks:

        block = block.strip()

        if not block:

            continue

        # Ignore document heading before Job 1.

        if not re.search(
            r"(?mi)^Job\s+\d+\s*:",
            block
        ):

            continue

        cleaned_blocks.append(block)

    return cleaned_blocks


# ==========================================================
# EXTRACT TOP LEVEL JOB NUMBER
# ==========================================================

def extract_job_number(block: str) -> str:

    match = re.search(
        r"(?mi)^Job\s+(\d+)\s*:",
        block
    )

    if match:

        return match.group(1).strip()

    return ""


# ==========================================================
# EXTRACT TITLE
# ==========================================================

def extract_title(block: str) -> str:

    # ------------------------------------------------------
    # Primary:
    #
    # Job 2: Data Engineer - Washington, DC...
    # ------------------------------------------------------

    match = re.search(
        r"(?mi)^Job\s+\d+\s*:\s*(.+?)\s*$",
        block
    )

    if match:

        return match.group(1).strip()

    # ------------------------------------------------------
    # Fallback:
    #
    # Title:
    # Data Engineer
    # ------------------------------------------------------

    match = re.search(
        r"(?mi)^Title:\s*\n\s*(.+?)\s*$",
        block
    )

    if match:

        return match.group(1).strip()

    return ""


# ==========================================================
# EXTRACT JOB URL
# ==========================================================

def extract_job_url(block: str) -> str:

    patterns = [

        # Primary exporter format
        r"(?mi)^Job URL:\s*(https?://\S+)",

        # Secondary embedded format
        r"(?mi)^URL:\s*\n\s*(https?://\S+)",

        # Generic Dice fallback
        r"(https?://(?:www\.)?dice\.com/job-detail/[^\s]+)"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            block
        )

        if match:

            url = match.group(1)

            url = url.rstrip(
                ").,;"
            )

            return url

    return ""


# ==========================================================
# EXTRACT ATS / DICE JOB ID
# ==========================================================

def extract_job_id(block: str) -> str:
    """
    The ATS Excel Job ID corresponds to values such as:

        Dice Id: 10429554
        Dice Id: RTL953449
        Dice Id: 91124817

    Therefore Dice Id is the strongest ID source.
    """

    patterns = [

        r"(?i)\bDice\s+Id:\s*([A-Za-z0-9_-]+)",

        r"(?im)^\s*Job\s*ID\s*[:\-]\s*([A-Za-z0-9_-]+)\s*$",

        r"(?im)^\s*Position\s*ID\s*[:\-]\s*([A-Za-z0-9_-]+)\s*$",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            block
        )

        if match:

            return match.group(1).strip()

    return ""


# ==========================================================
# BUILD JD INDEX
# ==========================================================

def build_jd_index(path: Path) -> list[dict[str, str]]:

    text = read_document_text(
        path
    )

    blocks = split_job_blocks(
        text
    )

    entries: list[dict[str, str]] = []

    for block in blocks:

        entry = {

            "job_number":
                extract_job_number(
                    block
                ),

            "job_id":
                extract_job_id(
                    block
                ),

            "url":
                extract_job_url(
                    block
                ),

            "title":
                extract_title(
                    block
                ),

            "description":
                block,
        }

        entries.append(
            entry
        )

    return entries


# ==========================================================
# NORMALIZE TITLE
# ==========================================================

def normalize_title(title: str) -> str:

    if not title:

        return ""

    title = str(title).lower()

    title = re.sub(
        r"\s+",
        " ",
        title
    )

    title = re.sub(
        r"[^\w\s]",
        " ",
        title
    )

    title = re.sub(
        r"\s+",
        " ",
        title
    )

    return title.strip()


# ==========================================================
# FIND JOB DESCRIPTION
# ==========================================================

def find_job_description(
    entries: list[dict[str, str]],
    job_id: str,
    job_url: str,
    job_name: str,
) -> str | None:
    """
    IMPORTANT:

    This argument order exactly matches main.py:

        find_job_description(
            jd_entries,
            job.job_id,
            job.job_url,
            job.job_name
        )
    """

    normalized_id = str(
        job_id or ""
    ).strip().lower()

    normalized_url = normalize_url(
        job_url
    )

    normalized_name = normalize_title(
        job_name
    )

    # ======================================================
    # 1. EXACT JOB ID MATCH
    # ======================================================

    if normalized_id:

        for entry in entries:

            entry_id = str(
                entry.get(
                    "job_id",
                    ""
                )
                or ""
            ).strip().lower()

            if (
                entry_id
                and
                entry_id == normalized_id
            ):

                return entry.get(
                    "description",
                    ""
                )

    # ======================================================
    # 2. EXACT URL MATCH
    # ======================================================

    if normalized_url:

        for entry in entries:

            entry_url = normalize_url(

                entry.get(
                    "url",
                    ""
                )
            )

            if (
                entry_url
                and
                entry_url == normalized_url
            ):

                return entry.get(
                    "description",
                    ""
                )

    # ======================================================
    # 3. JOB ID CONTAINED IN JD
    # ======================================================

    if normalized_id:

        for entry in entries:

            description = str(
                entry.get(
                    "description",
                    ""
                )
                or ""
            )

            if (
                normalized_id
                in description.lower()
            ):

                return description

    # ======================================================
    # 4. EXACT TITLE MATCH
    # ======================================================

    if normalized_name:

        for entry in entries:

            entry_title = normalize_title(

                entry.get(
                    "title",
                    ""
                )
            )

            if (
                entry_title
                and
                entry_title == normalized_name
            ):

                return entry.get(
                    "description",
                    ""
                )

    # ======================================================
    # 5. CONSERVATIVE PARTIAL TITLE MATCH
    # ======================================================

    if normalized_name:

        for entry in entries:

            entry_title = normalize_title(

                entry.get(
                    "title",
                    ""
                )
            )

            if not entry_title:

                continue

            # Require a reasonably substantial title
            # before doing containment matching.

            if (
                len(entry_title) >= 12
                and
                (
                    entry_title in normalized_name
                    or
                    normalized_name in entry_title
                )
            ):

                return entry.get(
                    "description",
                    ""
                )

    # ======================================================
    # NOTHING MATCHED
    # ======================================================

    return None