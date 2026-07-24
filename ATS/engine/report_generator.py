# ==========================================================
# ATS V4.0
# Excel Report Generator
# File: ATS/engine/report_generator.py
# ==========================================================

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ==========================================================
# REPORT GENERATOR
# ==========================================================

class ATSReportGenerator:

    def __init__(
        self,
        output_directory: Optional[str] = None,
    ) -> None:

        if output_directory:

            self.output_directory = Path(
                output_directory
            )

        else:

            self.output_directory = Path.cwd()

        self.output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        self.header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        self.header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        self.subheader_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        self.subheader_font = Font(
            bold=True,
        )

        self.apply_fill = PatternFill(
            fill_type="solid",
            fgColor="C6EFCE",
        )

        self.strong_match_fill = PatternFill(
            fill_type="solid",
            fgColor="E2F0D9",
        )

        self.review_fill = PatternFill(
            fill_type="solid",
            fgColor="FFF2CC",
        )

        self.low_match_fill = PatternFill(
            fill_type="solid",
            fgColor="FCE4D6",
        )

        self.error_fill = PatternFill(
            fill_type="solid",
            fgColor="FFC7CE",
        )

    # ======================================================
    # PUBLIC METHOD
    # ======================================================

    def generate(
        self,
        ranking_results: Iterable[Any],
        output_file: str,
    ) -> str:

        ranking_results = list(
            ranking_results or []
        )

        output_path = self._resolve_output_path(
            output_file
        )

        workbook = Workbook()

        default_sheet = workbook.active

        workbook.remove(
            default_sheet
        )

        summary_sheet = workbook.create_sheet(
            "Summary"
        )

        rankings_sheet = workbook.create_sheet(
            "Resume Rankings"
        )

        breakdown_sheet = workbook.create_sheet(
            "Detailed Breakdown"
        )

        missing_skills_sheet = workbook.create_sheet(
            "Missing Skills"
        )

        skill_frequency_sheet = workbook.create_sheet(
            "Skill Frequency"
        )

        self._write_summary_sheet(
            summary_sheet,
            ranking_results,
        )

        self._write_rankings_sheet(
            rankings_sheet,
            ranking_results,
        )

        self._write_breakdown_sheet(
            breakdown_sheet,
            ranking_results,
        )

        self._write_missing_skills_sheet(
            missing_skills_sheet,
            ranking_results,
        )

        self._write_skill_frequency_sheet(
            skill_frequency_sheet,
            ranking_results,
        )

        workbook.save(
            output_path
        )

        return str(output_path)

    # ======================================================
    # SUMMARY SHEET
    # ======================================================

    def _write_summary_sheet(
        self,
        worksheet: Worksheet,
        ranking_results: List[Any],
    ) -> None:

        headers = [
            "Job ID",
            "Job Name",
            "Job URL",
            "Best Resume",
            "Best ATS Score",
            "Recommendation",
            "Matched Required Skills",
            "Missing Required Skills",
            "Evidence Summary",
            "Error",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        row_number = 2

        for ranking in ranking_results:

            rankings = list(
                getattr(
                    ranking,
                    "rankings",
                    [],
                )
                or []
            )

            best_match = (
                rankings[0]
                if rankings
                else None
            )

            worksheet.cell(
                row=row_number,
                column=1,
                value=self._safe_string(
                    getattr(
                        ranking,
                        "job_id",
                        "",
                    )
                ),
            )

            worksheet.cell(
                row=row_number,
                column=2,
                value=self._safe_string(
                    getattr(
                        ranking,
                        "job_name",
                        "",
                    )
                ),
            )

            worksheet.cell(
                row=row_number,
                column=3,
                value=self._safe_string(
                    getattr(
                        ranking,
                        "job_url",
                        "",
                    )
                ),
            )

            worksheet.cell(
                row=row_number,
                column=4,
                value=self._safe_string(
                    getattr(
                        ranking,
                        "best_resume",
                        "",
                    )
                ),
            )

            worksheet.cell(
                row=row_number,
                column=5,
                value=self._safe_number(
                    getattr(
                        ranking,
                        "best_score",
                        0.0,
                    )
                ),
            )

            worksheet.cell(
                row=row_number,
                column=6,
                value=self._safe_string(
                    getattr(
                        ranking,
                        "recommendation",
                        "",
                    )
                ),
            )

            if best_match:

                worksheet.cell(
                    row=row_number,
                    column=7,
                    value=self._join_values(
                        getattr(
                            best_match,
                            "matched_required_skills",
                            [],
                        )
                    ),
                )

                worksheet.cell(
                    row=row_number,
                    column=8,
                    value=self._join_values(
                        getattr(
                            best_match,
                            "missing_required_skills",
                            [],
                        )
                    ),
                )

                worksheet.cell(
                    row=row_number, column=9,
                    value=self._safe_string(getattr(best_match, "evidence_summary", "")),
                )
                worksheet.cell(
                    row=row_number, column=10,
                    value=self._safe_string(getattr(best_match, "error", "")),
                )

            self._apply_recommendation_fill(
                worksheet=worksheet,
                row=row_number,
                column=6,
                recommendation=getattr(
                    ranking,
                    "recommendation",
                    "",
                ),
            )

            self._apply_url_hyperlink(
                worksheet=worksheet,
                row=row_number,
                column=3,
            )

            row_number += 1

        self._finalize_sheet(
            worksheet,
            freeze_cell="A2",
            auto_filter=True,
        )

    # ======================================================
    # RESUME RANKINGS SHEET
    # ======================================================

    def _write_rankings_sheet(
        self,
        worksheet: Worksheet,
        ranking_results: List[Any],
    ) -> None:

        headers = [
            "Job ID",
            "Job Name",
            "Job URL",
            "Rank",
            "Resume",
            "Final ATS Score",
            "Recommendation",
            "Required Skill Score",
            "Preferred Skill Score",
            "Semantic Score",
            "BM25 Score",
            "Role Score",
            "Experience Score",
            "Domain Score",
            "Certification Score",
            "Weighted Score",
            "Total Penalty",
            "Resume Role",
            "Job Role",
            "Resume Domain",
            "Job Domain",
            "Resume Experience",
            "Required Experience",
            "Matched Required Skills",
            "Missing Required Skills",
            "Matched Preferred Skills",
            "Missing Preferred Skills",
            "Matched Certifications",
            "Missing Certifications",
            "Target Platform",
            "Platform Years",
            "Project Count",
            "Strong Projects",
            "Demonstrated Required",
            "Listed/Weak Required",
            "Evidence Summary",
            "Warnings",
            "Error",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        row_number = 2

        for ranking in ranking_results:

            match_results = list(
                getattr(
                    ranking,
                    "rankings",
                    [],
                )
                or []
            )

            for rank_position, match in enumerate(
                match_results,
                start=1,
            ):

                scores = getattr(
                    match,
                    "scores",
                    None,
                )

                values = [
                    getattr(
                        match,
                        "job_id",
                        getattr(
                            ranking,
                            "job_id",
                            "",
                        ),
                    ),
                    getattr(
                        match,
                        "job_name",
                        getattr(
                            ranking,
                            "job_name",
                            "",
                        ),
                    ),
                    getattr(
                        match,
                        "job_url",
                        getattr(
                            ranking,
                            "job_url",
                            "",
                        ),
                    ),
                    rank_position,
                    getattr(
                        match,
                        "resume_name",
                        "",
                    ),
                    getattr(
                        match,
                        "final_score",
                        0.0,
                    ),
                    getattr(
                        match,
                        "recommendation",
                        "",
                    ),
                    self._score_value(
                        scores,
                        "required_skill_score",
                    ),
                    self._score_value(
                        scores,
                        "preferred_skill_score",
                    ),
                    self._score_value(
                        scores,
                        "semantic_score",
                    ),
                    self._score_value(
                        scores,
                        "bm25_score",
                    ),
                    self._score_value(
                        scores,
                        "role_score",
                    ),
                    self._score_value(
                        scores,
                        "experience_score",
                    ),
                    self._score_value(
                        scores,
                        "domain_score",
                    ),
                    self._score_value(
                        scores,
                        "certification_score",
                    ),
                    self._score_value(
                        scores,
                        "weighted_score",
                    ),
                    self._score_value(
                        scores,
                        "total_penalty",
                    ),
                    getattr(
                        match,
                        "resume_role",
                        "",
                    ),
                    getattr(
                        match,
                        "job_role",
                        "",
                    ),
                    getattr(
                        match,
                        "resume_domain",
                        "",
                    ),
                    getattr(
                        match,
                        "job_domain",
                        "",
                    ),
                    getattr(
                        match,
                        "resume_experience",
                        0.0,
                    ),
                    getattr(
                        match,
                        "required_experience",
                        0.0,
                    ),
                    self._join_values(
                        getattr(
                            match,
                            "matched_required_skills",
                            [],
                        )
                    ),
                    self._join_values(
                        getattr(
                            match,
                            "missing_required_skills",
                            [],
                        )
                    ),
                    self._join_values(
                        getattr(
                            match,
                            "matched_preferred_skills",
                            [],
                        )
                    ),
                    self._join_values(
                        getattr(
                            match,
                            "missing_preferred_skills",
                            [],
                        )
                    ),
                    self._join_values(
                        getattr(
                            match,
                            "matched_certifications",
                            [],
                        )
                    ),
                    self._join_values(
                        getattr(
                            match,
                            "missing_certifications",
                            [],
                        )
                    ),
                    getattr(match, "target_platform", ""),
                    getattr(match, "platform_years", 0.0),
                    getattr(match, "project_count", 0),
                    getattr(match, "strong_project_count", 0),
                    self._join_values(getattr(match, "demonstrated_required_skills", [])),
                    self._join_values(getattr(match, "listed_required_skills", [])),
                    getattr(match, "evidence_summary", ""),
                    self._join_values(
                        getattr(
                            match,
                            "warnings",
                            [],
                        ),
                        separator=" | ",
                    ),
                    getattr(
                        match,
                        "error",
                        "",
                    ),
                ]

                for column_number, value in enumerate(
                    values,
                    start=1,
                ):

                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                        value=self._excel_value(
                            value
                        ),
                    )

                self._apply_recommendation_fill(
                    worksheet=worksheet,
                    row=row_number,
                    column=7,
                    recommendation=getattr(
                        match,
                        "recommendation",
                        "",
                    ),
                )

                self._apply_url_hyperlink(
                    worksheet=worksheet,
                    row=row_number,
                    column=3,
                )

                row_number += 1

        self._finalize_sheet(
            worksheet,
            freeze_cell="A2",
            auto_filter=True,
        )

    # ======================================================
    # DETAILED BREAKDOWN SHEET
    # ======================================================

    def _write_breakdown_sheet(
        self,
        worksheet: Worksheet,
        ranking_results: List[Any],
    ) -> None:

        headers = [
            "Job ID",
            "Job Name",
            "Resume",
            "Component",
            "Raw Score",
            "Weight",
            "Weighted Contribution",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        row_number = 2

        component_mapping = [
            ("Platform Evidence", "platform_score", "platform"),
            ("Required Skill Evidence", "required_skill_score", "required_skills"),
            ("Role Compatibility", "role_score", "role"),
            ("Relevant Experience", "experience_score", "experience"),
            ("Preferred Skill Evidence", "preferred_skill_score", "preferred_skills"),
            ("Achievement Relevance", "achievement_score", "achievement"),
            ("Project Complexity", "complexity_score", "complexity"),
            ("Semantic Match", "semantic_score", "semantic"),
            ("BM25 Match", "bm25_score", "bm25"),
            ("Domain Match", "domain_score", "domain"),
            ("Certification Match", "certification_score", "certification"),
        ]

        for ranking in ranking_results:

            for match in getattr(
                ranking,
                "rankings",
                [],
            ) or []:

                scores = getattr(
                    match,
                    "scores",
                    None,
                )

                engine_weights = getattr(
                    match,
                    "weights",
                    {},
                ) or {}

                for (
                    component_name,
                    score_attribute,
                    weight_key,
                ) in component_mapping:

                    raw_score = self._score_value(
                        scores,
                        score_attribute,
                    )

                    weight = self._safe_number(
                        engine_weights.get(
                            weight_key,
                            0.0,
                        )
                    )

                    if weight == 0.0:

                        weight = self._default_weight(
                            weight_key
                        )

                    weighted_contribution = (
                        raw_score
                        * weight
                    )

                    row_values = [
                        getattr(
                            match,
                            "job_id",
                            getattr(
                                ranking,
                                "job_id",
                                "",
                            ),
                        ),
                        getattr(
                            match,
                            "job_name",
                            getattr(
                                ranking,
                                "job_name",
                                "",
                            ),
                        ),
                        getattr(
                            match,
                            "resume_name",
                            "",
                        ),
                        component_name,
                        raw_score,
                        weight,
                        round(
                            weighted_contribution,
                            2,
                        ),
                    ]

                    self._write_row(
                        worksheet,
                        row_number,
                        row_values,
                    )

                    row_number += 1

                penalty_rows = [
                    (
                        "Missing Required Penalty",
                        self._score_value(
                            scores,
                            "missing_required_penalty",
                        ),
                    ),
                    (
                        "Role Mismatch Penalty",
                        self._score_value(
                            scores,
                            "role_mismatch_penalty",
                        ),
                    ),
                    (
                        "Experience Penalty",
                        self._score_value(
                            scores,
                            "experience_penalty",
                        ),
                    ),
                    (
                        "Total Penalty",
                        self._score_value(
                            scores,
                            "total_penalty",
                        ),
                    ),
                    (
                        "Final ATS Score",
                        getattr(
                            match,
                            "final_score",
                            0.0,
                        ),
                    ),
                ]

                for component_name, score in penalty_rows:

                    self._write_row(
                        worksheet,
                        row_number,
                        [
                            getattr(
                                match,
                                "job_id",
                                getattr(
                                    ranking,
                                    "job_id",
                                    "",
                                ),
                            ),
                            getattr(
                                match,
                                "job_name",
                                getattr(
                                    ranking,
                                    "job_name",
                                    "",
                                ),
                            ),
                            getattr(
                                match,
                                "resume_name",
                                "",
                            ),
                            component_name,
                            self._safe_number(
                                score
                            ),
                            "",
                            "",
                        ],
                    )

                    row_number += 1

        self._finalize_sheet(
            worksheet,
            freeze_cell="A2",
            auto_filter=True,
        )

    # ======================================================
    # MISSING SKILLS SHEET
    # ======================================================

    def _write_missing_skills_sheet(
        self,
        worksheet: Worksheet,
        ranking_results: List[Any],
    ) -> None:

        headers = [
            "Job ID",
            "Job Name",
            "Resume",
            "Skill Type",
            "Missing Skill",
            "Final ATS Score",
            "Recommendation",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        row_number = 2

        for ranking in ranking_results:

            for match in getattr(
                ranking,
                "rankings",
                [],
            ) or []:

                missing_groups = [
                    (
                        "Required",
                        getattr(
                            match,
                            "missing_required_skills",
                            [],
                        ),
                    ),
                    (
                        "Preferred",
                        getattr(
                            match,
                            "missing_preferred_skills",
                            [],
                        ),
                    ),
                    (
                        "Certification",
                        getattr(
                            match,
                            "missing_certifications",
                            [],
                        ),
                    ),
                ]

                for skill_type, missing_items in missing_groups:

                    for missing_item in missing_items or []:

                        self._write_row(
                            worksheet,
                            row_number,
                            [
                                getattr(
                                    match,
                                    "job_id",
                                    getattr(
                                        ranking,
                                        "job_id",
                                        "",
                                    ),
                                ),
                                getattr(
                                    match,
                                    "job_name",
                                    getattr(
                                        ranking,
                                        "job_name",
                                        "",
                                    ),
                                ),
                                getattr(
                                    match,
                                    "resume_name",
                                    "",
                                ),
                                skill_type,
                                missing_item,
                                getattr(
                                    match,
                                    "final_score",
                                    0.0,
                                ),
                                getattr(
                                    match,
                                    "recommendation",
                                    "",
                                ),
                            ],
                        )

                        self._apply_recommendation_fill(
                            worksheet=worksheet,
                            row=row_number,
                            column=7,
                            recommendation=getattr(
                                match,
                                "recommendation",
                                "",
                            ),
                        )

                        row_number += 1

        self._finalize_sheet(
            worksheet,
            freeze_cell="A2",
            auto_filter=True,
        )

    # ======================================================
    # SKILL FREQUENCY SHEET
    # ======================================================

    def _write_skill_frequency_sheet(
        self,
        worksheet: Worksheet,
        ranking_results: List[Any],
    ) -> None:

        headers = [
            "Skill",
            "Required Frequency",
            "Preferred Frequency",
            "Total Job Frequency",
            "Matched Resume Frequency",
            "Missing Resume Frequency",
        ]

        self._write_headers(
            worksheet,
            headers,
        )

        required_counter: Counter = Counter()
        preferred_counter: Counter = Counter()
        matched_counter: Counter = Counter()
        missing_counter: Counter = Counter()

        job_skill_tracker: Dict[str, set] = {}

        for ranking in ranking_results:

            job_key = self._safe_string(
                getattr(
                    ranking,
                    "job_id",
                    "",
                )
            )

            if not job_key:

                job_key = self._safe_string(
                    getattr(
                        ranking,
                        "job_name",
                        "",
                    )
                )

            rankings = list(
                getattr(
                    ranking,
                    "rankings",
                    [],
                )
                or []
            )

            if not rankings:

                continue

            representative_match = rankings[0]

            required_skills = set(
                self._normalize_values(
                    list(
                        getattr(
                            representative_match,
                            "matched_required_skills",
                            [],
                        )
                        or []
                    )
                    +
                    list(
                        getattr(
                            representative_match,
                            "missing_required_skills",
                            [],
                        )
                        or []
                    )
                )
            )

            preferred_skills = set(
                self._normalize_values(
                    list(
                        getattr(
                            representative_match,
                            "matched_preferred_skills",
                            [],
                        )
                        or []
                    )
                    +
                    list(
                        getattr(
                            representative_match,
                            "missing_preferred_skills",
                            [],
                        )
                        or []
                    )
                )
            )

            required_counter.update(
                required_skills
            )

            preferred_counter.update(
                preferred_skills
            )

            job_skill_tracker[job_key] = (
                required_skills
                | preferred_skills
            )

            for match in rankings:

                matched_counter.update(
                    self._normalize_values(
                        list(
                            getattr(
                                match,
                                "matched_required_skills",
                                [],
                            )
                            or []
                        )
                        +
                        list(
                            getattr(
                                match,
                                "matched_preferred_skills",
                                [],
                            )
                            or []
                        )
                    )
                )

                missing_counter.update(
                    self._normalize_values(
                        list(
                            getattr(
                                match,
                                "missing_required_skills",
                                [],
                            )
                            or []
                        )
                        +
                        list(
                            getattr(
                                match,
                                "missing_preferred_skills",
                                [],
                            )
                            or []
                        )
                    )
                )

        all_skills = set(
            required_counter
        )

        all_skills.update(
            preferred_counter
        )

        all_skills.update(
            matched_counter
        )

        all_skills.update(
            missing_counter
        )

        total_job_frequency = Counter()

        for job_skills in job_skill_tracker.values():

            total_job_frequency.update(
                job_skills
            )

        sorted_skills = sorted(
            all_skills,
            key=lambda skill: (
                -total_job_frequency.get(
                    skill,
                    0,
                ),
                -required_counter.get(
                    skill,
                    0,
                ),
                skill,
            ),
        )

        row_number = 2

        for skill in sorted_skills:

            self._write_row(
                worksheet,
                row_number,
                [
                    skill,
                    required_counter.get(
                        skill,
                        0,
                    ),
                    preferred_counter.get(
                        skill,
                        0,
                    ),
                    total_job_frequency.get(
                        skill,
                        0,
                    ),
                    matched_counter.get(
                        skill,
                        0,
                    ),
                    missing_counter.get(
                        skill,
                        0,
                    ),
                ],
            )

            row_number += 1

        self._finalize_sheet(
            worksheet,
            freeze_cell="A2",
            auto_filter=True,
        )

    # ======================================================
    # STYLING HELPERS
    # ======================================================

    def _write_headers(
        self,
        worksheet: Worksheet,
        headers: List[str],
    ) -> None:

        for column_number, header in enumerate(
            headers,
            start=1,
        ):

            cell = worksheet.cell(
                row=1,
                column=column_number,
                value=header,
            )

            cell.fill = self.header_fill
            cell.font = self.header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        worksheet.row_dimensions[1].height = 30

    @staticmethod
    def _write_row(
        worksheet: Worksheet,
        row_number: int,
        values: List[Any],
    ) -> None:

        for column_number, value in enumerate(
            values,
            start=1,
        ):

            worksheet.cell(
                row=row_number,
                column=column_number,
                value=ATSReportGenerator._excel_value(
                    value
                ),
            )

    def _finalize_sheet(
        self,
        worksheet: Worksheet,
        freeze_cell: str,
        auto_filter: bool,
    ) -> None:

        worksheet.freeze_panes = freeze_cell

        if (
            auto_filter
            and worksheet.max_row >= 1
            and worksheet.max_column >= 1
        ):

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

        for row in worksheet.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        self._auto_size_columns(
            worksheet
        )

    @staticmethod
    def _auto_size_columns(
        worksheet: Worksheet,
    ) -> None:

        for column_cells in worksheet.columns:

            column_index = column_cells[0].column

            column_letter = get_column_letter(
                column_index
            )

            max_length = 0

            for cell in column_cells:

                value = cell.value

                if value is None:

                    continue

                value_length = len(
                    str(value)
                )

                if value_length > max_length:

                    max_length = value_length

            adjusted_width = min(
                max(
                    max_length + 2,
                    12,
                ),
                50,
            )

            worksheet.column_dimensions[
                column_letter
            ].width = adjusted_width

    def _apply_recommendation_fill(
        self,
        worksheet: Worksheet,
        row: int,
        column: int,
        recommendation: Any,
    ) -> None:

        recommendation_text = self._safe_string(
            recommendation
        ).upper()

        cell = worksheet.cell(
            row=row,
            column=column,
        )

        if recommendation_text == "APPLY":

            cell.fill = self.apply_fill

        elif recommendation_text == "STRONG MATCH":

            cell.fill = self.strong_match_fill

        elif recommendation_text == "REVIEW":

            cell.fill = self.review_fill

        elif recommendation_text == "ERROR":

            cell.fill = self.error_fill

        else:

            cell.fill = self.low_match_fill

    @staticmethod
    def _apply_url_hyperlink(
        worksheet: Worksheet,
        row: int,
        column: int,
    ) -> None:

        cell = worksheet.cell(
            row=row,
            column=column,
        )

        url = str(
            cell.value or ""
        ).strip()

        if not url.lower().startswith(
            (
                "http://",
                "https://",
            )
        ):

            return

        cell.hyperlink = url

        cell.style = "Hyperlink"

    # ======================================================
    # VALUE HELPERS
    # ======================================================

    def _resolve_output_path(
        self,
        output_file: str,
    ) -> Path:

        output_path = Path(
            output_file
        )

        if output_path.suffix.lower() != ".xlsx":

            output_path = output_path.with_suffix(
                ".xlsx"
            )

        if not output_path.is_absolute():

            output_path = (
                self.output_directory
                / output_path
            )

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        return output_path

    @staticmethod
    def _score_value(
        scores: Any,
        attribute: str,
    ) -> float:

        if scores is None:

            return 0.0

        return ATSReportGenerator._safe_number(
            getattr(
                scores,
                attribute,
                0.0,
            )
        )

    @staticmethod
    def _safe_number(
        value: Any,
    ) -> float:

        try:

            return round(
                float(
                    value or 0.0
                ),
                2,
            )

        except (
            TypeError,
            ValueError,
        ):

            return 0.0

    @staticmethod
    def _safe_string(
        value: Any,
    ) -> str:

        if value is None:

            return ""

        return str(value).strip()

    @staticmethod
    def _join_values(
        values: Any,
        separator: str = ", ",
    ) -> str:

        if not values:

            return ""

        if isinstance(
            values,
            str,
        ):

            return values.strip()

        return separator.join(
            str(value).strip()
            for value in values
            if str(value).strip()
        )

    @staticmethod
    def _normalize_values(
        values: Iterable[Any],
    ) -> List[str]:

        return sorted(
            {
                str(value).strip().lower()
                for value in values
                if str(value).strip()
            }
        )

    @staticmethod
    def _excel_value(
        value: Any,
    ) -> Any:

        if value is None:

            return ""

        if isinstance(
            value,
            (
                list,
                tuple,
                set,
            ),
        ):

            return ", ".join(
                str(item)
                for item in value
            )

        if isinstance(
            value,
            dict,
        ):

            return ", ".join(
                f"{key}: {item_value}"
                for key, item_value in value.items()
            )

        return value

    @staticmethod
    def _default_weight(
        weight_key: str,
    ) -> float:

        default_weights = {
            "platform": 0.30,
            "role": 0.15,
            "required_skills": 0.22,
            "preferred_skills": 0.04,
            "experience": 0.13,
            "semantic": 0.025,
            "bm25": 0.025,
            "domain": 0.005,
            "certification": 0.005,
            "achievement": 0.05,
            "complexity": 0.05,
        }

        return default_weights.get(
            weight_key,
            0.0,
        )