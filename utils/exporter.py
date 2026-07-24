from openpyxl import Workbook

from docx import Document

from docx.shared import Pt


class Exporter:

    def __init__(self, folder):

        self.folder = folder

        self.links = []

        self.jobs = []

        self.failed = []

    def add_job(

        self,

        url,

        title,

        description

    ):

        self.links.append(url)

        self.jobs.append(

            (url, title, description)

        )

    def save(self):

        wb = Workbook()

        ws = wb.active

        ws.title = "Job Links"

        ws["A1"] = "Job Link"

        for i, url in enumerate(

            self.links,

            start=2

        ):

            ws.cell(i, 1).value = url

        wb.save(

            self.folder / "links.xlsx"

        )

        doc = Document()

        doc.styles["Normal"].font.name = "Calibri"

        doc.styles["Normal"].font.size = Pt(11)

        txt = open(

            self.folder / "job_descriptions.txt",

            "w",

            encoding="utf-8"

        )

        for i, (

            url,

            title,

            desc

        ) in enumerate(

            self.jobs,

            start=1

        ):

            doc.add_heading(

                f"Job {i}",

                1

            )

            doc.add_paragraph(

                f"URL:\n{url}"

            )

            doc.add_paragraph(

                f"Title:\n{title}"

            )

            doc.add_heading(

                "Job Description",

                2

            )

            doc.add_paragraph(desc)

            doc.add_page_break()

            txt.write(

                f"Job {i}\n"

            )

            txt.write(

                f"URL:\n{url}\n"

            )

            txt.write(

                f"Title:\n{title}\n"

            )

            txt.write(desc)

            txt.write(

                "\n"+"="*80+"\n\n"

            )

        txt.close()

        doc.save(

            self.folder /

            "job_descriptions.docx"

        )

        with open(

            self.folder /

            "failed_urls.txt",

            "w",

            encoding="utf-8"

        ) as f:

            f.write(

                "\n".join(self.failed)

            )