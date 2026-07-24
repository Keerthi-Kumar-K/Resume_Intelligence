import os
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import (
    parse_qsl,
    urlencode,
    urlsplit,
    urlunsplit,
)

import pandas as pd
from docx import Document
from docx.shared import Pt


# ==============================================================
# CONFIGURATION
# ==============================================================

DATA_ROLE_KEYWORDS = [

    "data engineer",
    "senior data engineer",
    "lead data engineer",
    "principal data engineer",
    "staff data engineer",
    "big data engineer",
    "cloud data engineer",
    "azure data engineer",
    "aws data engineer",
    "gcp data engineer",

    "analytics engineer",
    "data platform engineer",
    "data infrastructure engineer",
    "data integration engineer",
    "streaming data engineer",
    "real time data engineer",
    "data pipeline engineer",

    "etl engineer",
    "elt engineer",
    "etl developer",
    "etl lead",
    "etl architect",
    "etl consultant",
    "data integration developer",
    "integration engineer",

    "data warehouse developer",
    "data warehouse engineer",
    "dwh developer",
    "dwh engineer",
    "warehouse developer",
    "warehouse engineer",

    "database developer",
    "database engineer",
    "database architect",
    "database analyst",

    "sql developer",
    "sql engineer",
    "pl/sql developer",
    "t-sql developer",

    "business intelligence",
    "bi developer",
    "bi engineer",
    "bi analyst",

    "power bi",
    "tableau",
    "looker",
    "qlik",
    "cognos",
    "microstrategy",
    "obiee",

    "data analyst",
    "analytics analyst",
    "analytics consultant",
    "report developer",
    "reporting developer",
    "reporting analyst",
    "insights analyst",
    "product analyst",

    "data scientist",
    "machine learning engineer",
    "ml engineer",
    "ai engineer",
    "genai engineer",
    "llm engineer",
    "deep learning engineer",
    "applied scientist",
    "research engineer",

    "data architect",
    "enterprise data architect",
    "cloud data architect",
    "analytics architect",
    "lakehouse architect",

    "data modeler",
    "data modelling",
    "data modeling",

    "master data",
    "mdm",
    "data governance",
    "data steward",
    "data quality",

    "metadata engineer",
    "metadata analyst",

    "snowflake",
    "databricks",
    "spark",
    "pyspark",
    "hadoop",
    "kafka",
    "airflow",
    "dbt",
    "dagster",
    "prefect",
    "apache beam",

    "microsoft fabric",
    "azure fabric",
    "synapse",
    "redshift",
    "bigquery",
    "azure data factory",
    "aws glue",

    "informatica",
    "ab initio",
    "talend",
    "datastage",
    "ssis",
    "matillion",
    "fivetran",
    "dataiku",

    "data warehouse",
    "lakehouse",
    "data pipeline",
    "data platform",
    "big data",

    "business analyst",
    "business systems analyst"

]


REJECT_TITLE_KEYWORDS = [

    "frontend",
    "front end",
    "backend",
    "back end",
    "full stack",
    "react developer",
    "angular developer",
    "vue developer",
    "ios developer",
    "android developer",
    "mobile developer",
    "web developer",

    "java developer",
    ".net developer",
    "dotnet developer",
    "c# developer",
    "c++ developer",
    "golang developer",
    "ruby developer",
    "php developer",
    "node developer",
    "software developer",
    "application developer",

    "devops engineer",
    "site reliability engineer",
    "sre engineer",

    "network engineer",
    "network administrator",
    "systems engineer",
    "linux engineer",
    "windows engineer",
    "storage engineer",

    "security engineer",
    "cyber security",
    "cybersecurity",
    "information security",
    "soc analyst",
    "penetration tester",

    "qa engineer",
    "quality assurance",
    "automation tester",
    "test engineer",
    "manual tester",
    "selenium tester",

    "salesforce",
    "servicenow",
    "service now",

    "guidewire",
    "policycenter",
    "claimcenter",
    "billingcenter",

    "sap fico",
    "sap basis",
    "sap abap",
    "sap mm",
    "sap sd",

    "oracle ebs",
    "oracle fusion",

    "embedded",
    "firmware",
    "fpga",
    "asic",

    "electrical engineer",
    "mechanical engineer",

    "desktop support",
    "help desk",
    "technical support",
    "it support",
    "field technician",

    "project manager",
    "program manager",
    "product manager",
    "scrum master",
    "delivery manager",
    "engineering manager",

    "director",
    "vice president",
    "chief",

    "recruiter",
    "talent acquisition",
    "human resources",

    "graphic designer",
    "ux designer",
    "ui designer",

    "accountant",
    "auditor"

]


TRACKING_QUERY_KEYS = {

    "utm_source",
    "utm_medium",
    "utm_campaign",
    "utm_term",
    "utm_content",
    "source",
    "src",
    "tracking",
    "ref",
    "referrer",
    "gh_src",
    "lever-source"

}


EXPERIENCE_REGEX = re.compile(

    r"(?ix)\b(?:"
    r"(?:at\s+least|min(?:imum)?(?:\s+of)?|minimum|required|requires?|"
    r"must\s+have|over|more\s+than|greater\s+than|typically)?\s*"
    r"(?:10|1[1-9]|[2-9]\d)\s*(?:\+|plus)?\s*"
    r"(?:years?|yrs?)"
    r"|(?:ten|eleven|twelve|thirteen|fourteen|fifteen|sixteen|"
    r"seventeen|eighteen|nineteen|twenty)\s*(?:\+|plus)?\s*"
    r"(?:years?|yrs?)"
    r")\b"

)


URL_REGEX = re.compile(

    r"https?://[^\s<>\"]+",
    re.I

)


STRUCTURED_START = re.compile(

    r"(?im)^\s*(?:={5,}\s*)?JOB\s+START(?:\s*={5,})?\s*$"

)

STRUCTURED_END = re.compile(

    r"(?im)^\s*(?:={5,}\s*)?JOB\s+END(?:\s*={5,})?\s*$"

)

PROCESSING_MARKER = re.compile(

    r"(?im)^\s*Processing\s*:\s*.*$"

)

SEPARATOR_MARKER = re.compile(

    r"(?m)^\s*={20,}\s*$"

)


JOB_NUMBER_MARKER = re.compile(

    r"(?m)^Job\s+\d+\s*$"

)


# ==============================================================
# BASIC UTILITIES
# ==============================================================

def clean_space(value):

    return re.sub(
        r"\s+",
        " ",
        str(value or "")
    ).strip()


def normalize_url(value):

    raw = str(value or "").strip()

    raw = raw.strip('"').strip("'")

    raw = raw.rstrip(".,);]}")

    if not raw:

        return ""

    if raw.lower() == "nan":

        return ""

    try:

        parts = urlsplit(raw)

        if not parts.scheme or not parts.netloc:

            return raw.lower().rstrip("/")

        host = parts.netloc.lower()

        if host.startswith("www."):

            host = host[4:]

        path = re.sub(
            r"/{2,}",
            "/",
            parts.path
        ).rstrip("/")

        query = [

            (k, v)

            for k, v in parse_qsl(
                parts.query,
                keep_blank_values=True
            )

            if k.lower() not in TRACKING_QUERY_KEYS

        ]

        query = urlencode(sorted(query))

        return urlunsplit(

            (

                parts.scheme.lower(),
                host,
                path,
                query,
                ""

            )

        )

    except Exception:

        return raw.lower().rstrip("/")


def keyword_in_text(keywords, text):

    text = text.lower()

    for keyword in keywords:

        if keyword in text:

            return keyword

    return ""


def first_matching_pattern(patterns, text):

    for pattern in patterns:

        match = pattern.search(text)

        if match:

            return clean_space(

                match.group(0)

            )

    return ""
    
# ==============================================================
# INPUT READERS
# ==============================================================

def read_docx(file_path):

    doc = Document(file_path)

    blocks = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:

            blocks.append(text)

    for table in doc.tables:

        for row in table.rows:

            cells = [

                clean_space(cell.text)

                for cell in row.cells

            ]

            row_text = " | ".join(

                cell

                for cell in cells

                if cell

            )

            if row_text:

                blocks.append(row_text)

    return "\n".join(blocks)


def read_txt(file_path):

    with open(

        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"

    ) as handle:

        return handle.read()


def read_jd_file(file_path):

    suffix = Path(file_path).suffix.lower()

    if suffix == ".docx":

        return read_docx(file_path)

    if suffix == ".txt":

        return read_txt(file_path)

    raise ValueError(

        "JD file must be .docx or .txt"

    )


# ==============================================================
# EXCEL COLUMN DISCOVERY
# ==============================================================

def find_column(

    df,
    candidates,
    contains=None

):

    normalized = {

        clean_space(col).lower(): col

        for col in df.columns

    }

    for candidate in candidates:

        if candidate.lower() in normalized:

            return normalized[candidate.lower()]

    if contains:

        for col in df.columns:

            lower = clean_space(col).lower()

            if any(

                token in lower

                for token in contains

            ):

                return col

    return None


def find_url_column(df):

    column = find_column(

        df,

        [

            "url",
            "job url",
            "job link",
            "link",
            "links"

        ],

        [

            "url",
            "link"

        ]

    )

    if column is None:

        raise ValueError(

            "Unable to locate URL column."

        )

    return column


def load_links(file_path):

    df = pd.read_excel(file_path)

    if df.empty:

        raise ValueError(

            "Links Excel is empty."

        )

    url_col = find_url_column(df)

    title_col = find_column(

        df,

        [

            "job title",
            "title",
            "job name"

        ],

        [

            "title"

        ]

    )

    company_col = find_column(

        df,

        [

            "company",
            "company name",
            "employer"

        ],

        [

            "company"

        ]

    )

    location_col = find_column(

        df,

        [

            "location",
            "job location"

        ],

        [

            "location"

        ]

    )

    df = df.copy()

    df["__normalized_url"] = (

        df[url_col]

        .map(normalize_url)

    )

    df["__row_order"] = range(

        len(df)

    )

    return df, {

        "url": url_col,
        "title": title_col,
        "company": company_col,
        "location": location_col

    }


# ==============================================================
# JOB PARSER
# ==============================================================

def split_jobs(raw_text):

    text = (

        raw_text

        .replace("\r\n", "\n")

        .replace("\r", "\n")

        .strip()

    )

    if not text:

        return []

    # ----------------------------------------------------------
    # FORMAT 1
    # JOB START / JOB END
    # ----------------------------------------------------------

    starts = list(

        STRUCTURED_START.finditer(text)

    )

    if starts:

        jobs = []

        pos = 0

        while True:

            start = STRUCTURED_START.search(

                text,
                pos

            )

            if not start:

                break

            end = STRUCTURED_END.search(

                text,
                start.end()

            )

            if end:

                chunk = text[

                    start.end():

                    end.start()

                ].strip()

                pos = end.end()

            else:

                chunk = text[

                    start.end():

                ].strip()

                pos = len(text)

            if chunk:

                jobs.append(chunk)

        return jobs

    # ----------------------------------------------------------
    # FORMAT 2
    # Dice Export
    # Job 1
    # Job 2
    # Job 3
    # ----------------------------------------------------------

    markers = list(

        JOB_NUMBER_MARKER.finditer(text)

    )

    if markers:

        jobs = []

        for i, marker in enumerate(markers):

            start = marker.start()

            end = (

                markers[i + 1].start()

                if i + 1 < len(markers)

                else len(text)

            )

            chunk = text[

                start:end

            ].strip()

            chunk = re.split(

                r"(?im)^Similar Jobs.*$",

                chunk,

                maxsplit=1

            )[0]

            chunk = re.split(

                r"(?im)^Technology Professionals.*$",

                chunk,

                maxsplit=1

            )[0]

            chunk = re.split(

                r"(?im)^Dice Id.*$",

                chunk,

                maxsplit=1

            )[0]

            chunk = re.sub(

                r"\n{3,}",

                "\n\n",

                chunk

            ).strip()

            if chunk:

                jobs.append(chunk)

        return jobs

    # ----------------------------------------------------------
    # FORMAT 3
    # Processing:
    # ----------------------------------------------------------

    processing = list(

        PROCESSING_MARKER.finditer(text)

    )

    if processing:

        jobs = []

        for i, marker in enumerate(processing):

            end = (

                processing[i + 1].start()

                if i + 1 < len(processing)

                else len(text)

            )

            chunk = text[

                marker.start():

                end

            ].strip()

            if chunk:

                jobs.append(chunk)

        return jobs

    # ----------------------------------------------------------
    # FORMAT 4
    # ========
    # ----------------------------------------------------------

    chunks = [

        c.strip()

        for c in SEPARATOR_MARKER.split(text)

        if c.strip()

    ]

    if len(chunks) > 1:

        return chunks

    return [

        text

    ]
    
# ==============================================================
# FIELD EXTRACTION
# ==============================================================

def extract_label(job, labels):

    for label in labels:

        pattern = re.compile(

            rf"(?im)^\s*{re.escape(label)}\s*[:\-]\s*(.+?)\s*$"

        )

        match = pattern.search(job)

        if match:

            return clean_space(

                match.group(1)

            )

    return ""


def extract_url(job):

    labeled = extract_label(

        job,

        [

            "Job URL",
            "URL",
            "Link",
            "Job Link"

        ]

    )

    if labeled:

        match = URL_REGEX.search(labeled)

        if match:

            return normalize_url(

                match.group(0)

            )

        return normalize_url(

            labeled

        )

    match = URL_REGEX.search(job)

    if match:

        return normalize_url(

            match.group(0)

        )

    return ""


def extract_title(job):

    title = extract_label(

        job,

        [

            "Job Title",
            "Title",
            "Job Name",
            "Position"

        ]

    )

    if title:

        return title

    processing = PROCESSING_MARKER.search(job)

    if processing:

        value = processing.group(0)

        value = value.split(":", 1)[1].strip()

        if value and not re.fullmatch(

            r"(?:job\s*)?\d+",

            value,

            re.I

        ):

            return clean_space(value)

    lines = [

        clean_space(line)

        for line in job.splitlines()

        if clean_space(line)

    ]

    ignore = (

        "job ",
        "url:",
        "job url:",
        "link:",
        "company:",
        "location:",
        "description:",
        "job description:",
        "processing:"

    )

    for line in lines:

        lower = line.lower()

        if lower.startswith(ignore):

            continue

        if URL_REGEX.fullmatch(line):

            continue

        if len(line) <= 180:

            return line

    return "Unknown"


def extract_company(job):

    return extract_label(

        job,

        [

            "Company",
            "Company Name",
            "Employer"

        ]

    )


def extract_location(job):

    return extract_label(

        job,

        [

            "Location",
            "Job Location"

        ]

    )


# ==============================================================
# BUILD JOBS
# ==============================================================

def build_jobs(

    raw_text,

    links_df,

    columns

):

    raw_jobs = split_jobs(

        raw_text

    )

    jobs = []

    excel_lookup = {

        row["__normalized_url"]: row

        for _, row in links_df.iterrows()

        if row["__normalized_url"]

    }

    for idx, chunk in enumerate(

        raw_jobs,

        start=1

    ):

        parsed_url = extract_url(chunk)

        excel_row = excel_lookup.get(

            parsed_url

        )

        if (

            excel_row is None

            and

            len(raw_jobs) == len(links_df)

        ):

            excel_row = links_df.iloc[idx - 1]

        if excel_row is not None:

            url = parsed_url or normalize_url(

                excel_row[columns["url"]]

            )

            title = extract_title(chunk)

            if (

                title == "Unknown"

                or

                not title

            ) and columns["title"]:

                title = clean_space(

                    excel_row[

                        columns["title"]

                    ]

                )

            company = extract_company(chunk)

            if (

                not company

                and

                columns["company"]

            ):

                company = clean_space(

                    excel_row[

                        columns["company"]

                    ]

                )

            location = extract_location(chunk)

            if (

                not location

                and

                columns["location"]

            ):

                location = clean_space(

                    excel_row[

                        columns["location"]

                    ]

                )

        else:

            url = parsed_url

            title = extract_title(chunk)

            company = extract_company(chunk)

            location = extract_location(chunk)

        jobs.append(

            {

                "index": idx,

                "title": title,

                "company": company,

                "location": location,

                "url": url,

                "text": chunk.strip()

            }

        )

    return jobs
    
# ==============================================================
# FILTERING
# ==============================================================

CITIZENSHIP_REGEXES = [

    re.compile(pattern, re.I)

    for pattern in [

        r"\b(?:u\.?s\.?|united states)\s+citizens?\s+only\b",

        r"\bmust\s+be\s+(?:a\s+)?(?:u\.?s\.?|united states)\s+citizen\b",

        r"\b(?:usc|u\.s\.c\.)\s*(?:only|required)\b",

        r"\bgreen\s*card\s*(?:only|required|holders?\s+only)\b",

        r"\bgc\s*(?:only|required)\b",

        r"\b(?:usc|u\.?s\.?\s+citizen)\s*(?:/|or|and)\s*(?:gc|green\s*card)\b",

        r"\bpermanent\s+residents?\s+only\b",

        r"\bno\s+(?:h-?1b|opt|cpt|visa(?:s)?)\b",

        r"\b(?:h-?1b|opt|cpt)\s+(?:candidates?|holders?)\s+(?:will\s+not\s+be|are\s+not)\s+considered\b",

        r"\bcannot\s+(?:provide|offer|support)\s+(?:visa\s+)?sponsorship\b",

        r"\bwill\s+not\s+(?:provide|offer|support)\s+(?:visa\s+)?sponsorship\b",

        r"\bno\s+(?:visa\s+)?sponsorship\b",

        r"\bwithout\s+(?:current\s+or\s+future\s+)?sponsorship\b",

        r"\bmust\s+not\s+require\s+(?:current\s+or\s+future\s+)?sponsorship\b"

    ]

]


CLEARANCE_REGEXES = [

    re.compile(pattern, re.I)

    for pattern in [

        r"\bactive\s+(?:security\s+)?clearance\b",

        r"\bsecurity\s+clearance\s+(?:is\s+)?required\b",

        r"\bmust\s+(?:hold|have|maintain|possess|obtain)\s+(?:an?\s+)?(?:active\s+)?(?:security\s+)?clearance\b",

        r"\b(?:top\s+secret|secret|confidential)\s+clearance\b",

        r"\bts\s*/\s*sci\b",

        r"\b(?:dod|department\s+of\s+defense|federal)\s+clearance\b",

        r"\bpublic\s+trust\s+(?:clearance\s+)?(?:is\s+)?required\b",

        r"\bability\s+to\s+obtain\s+(?:and\s+maintain\s+)?(?:a\s+)?(?:security\s+)?clearance\b"

    ]

]


def experience_reason(text):

    match = EXPERIENCE_REGEX.search(text)

    if not match:

        return ""

    context = text[

        max(0, match.start() - 100):

        min(len(text), match.end() + 100)

    ].lower()

    preferred_words = [

        "preferred",

        "nice to have",

        "desired",

        "plus",

        "bonus",

        "good to have"

    ]

    required_words = [

        "required",

        "must",

        "minimum",

        "at least",

        "need",

        "requires"

    ]

    if (

        any(word in context for word in preferred_words)

        and

        not any(word in context for word in required_words)

    ):

        return ""

    return clean_space(

        match.group(0)

    )


def citizenship_reason(text):

    return first_matching_pattern(

        CITIZENSHIP_REGEXES,

        text

    )


def clearance_reason(text):

    return first_matching_pattern(

        CLEARANCE_REGEXES,

        text

    )


def non_data_reason(

    title,

    text

):

    title = clean_space(title)

    lower_title = title.lower()

    positive = keyword_in_text(

        DATA_ROLE_KEYWORDS,

        lower_title

    )

    if positive:

        return ""

    negative = keyword_in_text(

        REJECT_TITLE_KEYWORDS,

        lower_title

    )

    if negative:

        return f"Irrelevant title keyword: {negative}"

    intro = " ".join(

        text.splitlines()[:20]

    ).lower()

    if keyword_in_text(

        DATA_ROLE_KEYWORDS,

        intro

    ):

        return ""

    return "No recognized data-role keyword"


def evaluate_job(job):

    reasons = []

    text = job["text"]

    exp = experience_reason(text)

    if exp:

        reasons.append(

            f"10+ years required: {exp}"

        )

    citizenship = citizenship_reason(text)

    if citizenship:

        reasons.append(

            f"Citizenship/Visa: {citizenship}"

        )

    clearance = clearance_reason(text)

    if clearance:

        reasons.append(

            f"Security Clearance: {clearance}"

        )

    non_data = non_data_reason(

        job["title"],

        text

    )

    if non_data:

        reasons.append(

            f"Non Data Role: {non_data}"

        )

    return reasons
    
# ==============================================================
# EXPORTERS
# ==============================================================

def write_filtered_docx(

    path,

    jobs

):

    doc = Document()

    doc.add_heading(

        "Filtered Job Descriptions",

        level=1

    )

    for index, job in enumerate(

        jobs,

        start=1

    ):

        doc.add_heading(

            f"Job {index}: {job['title']}",

            level=2

        )

        metadata = []

        if job["company"]:

            metadata.append(

                f"Company: {job['company']}"

            )

        if job["location"]:

            metadata.append(

                f"Location: {job['location']}"

            )

        if job["url"]:

            metadata.append(

                f"Job URL: {job['url']}"

            )

        if metadata:

            paragraph = doc.add_paragraph()

            paragraph.style.font.name = "Calibri"

            paragraph.style.font.size = Pt(10)

            paragraph.add_run(

                "\n".join(metadata)

            )

        body = doc.add_paragraph()

        body.style.font.name = "Calibri"

        body.style.font.size = Pt(10)

        body.add_run(

            job["text"]

        )

        if index < len(jobs):

            doc.add_page_break()

    doc.save(path)


def write_filtered_txt(

    path,

    jobs

):

    with open(

        path,

        "w",

        encoding="utf-8"

    ) as handle:

        for index, job in enumerate(

            jobs,

            start=1

        ):

            handle.write(

                "=" * 100 + "\n"

            )

            handle.write(

                f"JOB {index}\n"

            )

            handle.write(

                f"Job Title: {job['title']}\n"

            )

            if job["company"]:

                handle.write(

                    f"Company: {job['company']}\n"

                )

            if job["location"]:

                handle.write(

                    f"Location: {job['location']}\n"

                )

            if job["url"]:

                handle.write(

                    f"Job URL: {job['url']}\n"

                )

            handle.write(

                "-" * 100 + "\n"

            )

            handle.write(

                job["text"].strip()

            )

            handle.write(

                "\n\n"

            )


def build_final_links(

    links_df,

    accepted_jobs

):

    accepted_urls = {

        job["url"]

        for job in accepted_jobs

        if job["url"]

    }

    if accepted_urls:

        final_df = links_df[

            links_df["__normalized_url"].isin(

                accepted_urls

            )

        ].copy()

    elif len(accepted_jobs) == len(links_df):

        final_df = links_df.copy()

    else:

        final_df = links_df.iloc[0:0].copy()

    final_df = final_df.sort_values(

        "__row_order"

    )

    return final_df.drop(

        columns=[

            "__normalized_url",

            "__row_order"

        ],

        errors="ignore"

    )


def build_rejected_dataframe(

    rejected_jobs

):

    rows = []

    for job in rejected_jobs:

        rows.append(

            {

                "Job Number": job["index"],

                "Job Title": job["title"],

                "Company": job["company"],

                "Location": job["location"],

                "URL": job["url"],

                "Rejection Reasons": " | ".join(

                    job["reasons"]

                ),

                "Job Description": job["text"]

            }

        )

    return pd.DataFrame(

        rows

    )


def autosize_excel(

    path

):

    try:

        from openpyxl import load_workbook

        workbook = load_workbook(

            path

        )

        for sheet in workbook.worksheets:

            sheet.freeze_panes = "A2"

            sheet.auto_filter.ref = sheet.dimensions

            for column_cells in sheet.columns:

                letter = column_cells[0].column_letter

                max_length = max(

                    len(str(cell.value))

                    if cell.value is not None

                    else 0

                    for cell in column_cells

                )

                sheet.column_dimensions[

                    letter

                ].width = min(

                    max(

                        max_length + 2,

                        12

                    ),

                    60

                )

        workbook.save(path)

    except Exception:

        pass
        
# ==============================================================
# MAIN
# ==============================================================

def main():

    jd_file = input(

        "Enter JD Word/Text file: "

    ).strip().strip('"')

    link_file = input(

        "Enter Links Excel: "

    ).strip().strip('"')

    if not os.path.isfile(jd_file):

        raise FileNotFoundError(

            f"JD file not found: {jd_file}"

        )

    if not os.path.isfile(link_file):

        raise FileNotFoundError(

            f"Links Excel not found: {link_file}"

        )

    output_dir = os.path.join(

        os.path.dirname(

            os.path.abspath(jd_file)

        ),

        "Filtered"

    )

    os.makedirs(

        output_dir,

        exist_ok=True

    )

    timestamp = datetime.now().strftime(

        "%Y%m%d_%H%M%S"

    )

    filtered_docx = os.path.join(

        output_dir,

        f"Filtered_JDs_{timestamp}.docx"

    )

    filtered_txt = os.path.join(

        output_dir,

        f"Filtered_JDs_{timestamp}.txt"

    )

    final_links_file = os.path.join(

        output_dir,

        f"Final_Links_{timestamp}.xlsx"

    )

    rejected_file = os.path.join(

        output_dir,

        f"Rejected_JDs_{timestamp}.xlsx"

    )

    raw_text = read_jd_file(

        jd_file

    )

    links_df, columns = load_links(

        link_file

    )

    jobs = build_jobs(

        raw_text,

        links_df,

        columns

    )

    accepted_jobs = []

    rejected_jobs = []

    print()

    print("=" * 80)

    print("JD FILTER ENGINE")

    print("=" * 80)

    print(f"Parsed JDs   : {len(jobs)}")

    print(f"Excel Links  : {len(links_df)}")

    print()

    for job in jobs:

        reasons = evaluate_job(job)

        if reasons:

            job["reasons"] = reasons

            rejected_jobs.append(job)

            print(

                f"[{job['index']}/{len(jobs)}] REJECTED - {job['title']}"

            )

            for reason in reasons:

                print(

                    "   -",

                    reason

                )

        else:

            accepted_jobs.append(job)

            print(

                f"[{job['index']}/{len(jobs)}] ACCEPTED - {job['title']}"

            )

    final_links = build_final_links(

        links_df,

        accepted_jobs

    )

    rejected_df = build_rejected_dataframe(

        rejected_jobs

    )

    write_filtered_docx(

        filtered_docx,

        accepted_jobs

    )

    write_filtered_txt(

        filtered_txt,

        accepted_jobs

    )

    final_links.to_excel(

        final_links_file,

        index=False

    )

    rejected_df.to_excel(

        rejected_file,

        index=False

    )

    autosize_excel(

        final_links_file

    )

    autosize_excel(

        rejected_file

    )

    print()

    print("=" * 80)

    print("COMPLETED")

    print("=" * 80)

    print(f"Input JDs      : {len(jobs)}")

    print(f"Accepted JDs   : {len(accepted_jobs)}")

    print(f"Rejected JDs   : {len(rejected_jobs)}")

    print(f"Final Links    : {len(final_links)}")

    print()

    print("Generated Files")

    print("------------------------------")

    print(filtered_docx)

    print(filtered_txt)

    print(final_links_file)

    print(rejected_file)


if __name__ == "__main__":

    main()

# END OF 